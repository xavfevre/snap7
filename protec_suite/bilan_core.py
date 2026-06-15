#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Bilan kilométrique des déchets - PROTEC (moteur sans interface graphique)"""

import os
import json
import time
import math
from urllib import request as urlrequest, parse as urlparse

APP_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(APP_DIR, "config_bilan.json")
GEO_CACHE_PATH = os.path.join(APP_DIR, "cache_geocodage.json")
ROUTE_CACHE_PATH = os.path.join(APP_DIR, "cache_itineraires.json")

TD_API_URL = "https://api.trackdechets.beta.gouv.fr/"
BAN_URL = "https://api-adresse.data.gouv.fr/search/"
ORS_URL = "https://api.openrouteservice.org/v2/directions/driving-car"
OSRM_URL = "https://router.project-osrm.org/route/v1/driving/"

HAVERSINE_ROAD_FACTOR = 1.3

DEFAULT_CONFIG = {
    "td_api_token": "",
    "ors_api_key": "",
    "protec_siret": "",
    "protec_adresse": "",
    "round_trip": True,
    "categories": {
        "Emballages vides souilles plastique": ["150110"],
        "Huiles": ["1302"],
        "Cartons et papier": ["150101", "200101"],
        "Solvants": ["140603", "200113", "070104"],
        "Dechets hydrocarbures": ["1305", "160708", "130507"]
    }
}


def load_config():
    if os.path.exists(CONFIG_PATH):
        try:
            with open(CONFIG_PATH, "r", encoding="utf-8-sig") as f:
                cfg = json.load(f)
            for k, v in DEFAULT_CONFIG.items():
                cfg.setdefault(k, v)
            return cfg
        except Exception:
            pass
    return json.loads(json.dumps(DEFAULT_CONFIG))


def save_config(cfg):
    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


def _load_json(path):
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8-sig") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}


def _save_json(path, data):
    try:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)
    except Exception:
        pass


def http_json(url, method="GET", headers=None, payload=None, timeout=30):
    headers = headers or {}
    data = None
    if payload is not None:
        data = json.dumps(payload).encode("utf-8")
        headers.setdefault("Content-Type", "application/json")
    req = urlrequest.Request(url, data=data, headers=headers, method=method)
    with urlrequest.urlopen(req, timeout=timeout) as resp:
        raw = resp.read().decode("utf-8")
    return json.loads(raw)


class GeoService:
    def __init__(self, ors_key="", logger=print):
        self.ors_key = (ors_key or "").strip()
        self.log = logger
        self.geo_cache = _load_json(GEO_CACHE_PATH)
        self.route_cache = _load_json(ROUTE_CACHE_PATH)
        self._dirty_geo = False
        self._dirty_route = False

    def geocode(self, address):
        if not address:
            return None
        key = " ".join(address.lower().split())
        if key in self.geo_cache:
            v = self.geo_cache[key]
            return tuple(v) if v else None
        coords = None
        try:
            q = urlparse.urlencode({"q": address, "limit": 1})
            data = http_json(f"{BAN_URL}?{q}")
            feats = data.get("features") or []
            if feats:
                coords = tuple(feats[0]["geometry"]["coordinates"])
        except Exception as e:
            self.log(f"   ! géocodage échoué pour « {address} » : {e}")
        self.geo_cache[key] = list(coords) if coords else None
        self._dirty_geo = True
        time.sleep(0.05)
        return coords

    @staticmethod
    def _haversine_km(c1, c2):
        lon1, lat1 = c1
        lon2, lat2 = c2
        r = 6371.0
        p1, p2 = math.radians(lat1), math.radians(lat2)
        dphi = math.radians(lat2 - lat1)
        dlmb = math.radians(lon2 - lon1)
        a = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlmb / 2) ** 2
        return 2 * r * math.asin(math.sqrt(a))

    def _route_ors(self, c1, c2):
        if not self.ors_key:
            return None
        try:
            payload = {"coordinates": [[c1[0], c1[1]], [c2[0], c2[1]]]}
            headers = {"Authorization": self.ors_key}
            data = http_json(ORS_URL, method="POST", headers=headers, payload=payload)
            return data["routes"][0]["summary"]["distance"] / 1000.0
        except Exception:
            return None

    def _route_osrm(self, c1, c2):
        try:
            coords = f"{c1[0]},{c1[1]};{c2[0]},{c2[1]}"
            url = f"{OSRM_URL}{coords}?overview=false"
            data = http_json(url)
            if data.get("code") == "Ok" and data.get("routes"):
                return data["routes"][0]["distance"] / 1000.0
        except Exception:
            return None
        return None

    def road_km(self, c1, c2):
        if not c1 or not c2:
            return None, "n/a"
        rk = (round(c1[0], 5), round(c1[1], 5), round(c2[0], 5), round(c2[1], 5))
        ck = "|".join(str(x) for x in rk)
        if ck in self.route_cache:
            v = self.route_cache[ck]
            return v["km"], v["src"]
        km = self._route_ors(c1, c2)
        src = "ORS"
        if km is None:
            km = self._route_osrm(c1, c2)
            src = "OSRM"
        if km is None:
            km = self._haversine_km(c1, c2) * HAVERSINE_ROAD_FACTOR
            src = "estim."
        self.route_cache[ck] = {"km": round(km, 2), "src": src}
        self._dirty_route = True
        time.sleep(0.05)
        return round(km, 2), src

    def flush(self):
        if self._dirty_geo:
            _save_json(GEO_CACHE_PATH, self.geo_cache)
            self._dirty_geo = False
        if self._dirty_route:
            _save_json(ROUTE_CACHE_PATH, self.route_cache)
            self._dirty_route = False


FORMS_QUERY = """
query forms($siret: String, $first: Int, $cursorAfter: ID) {
  forms(siret: $siret, first: $first, cursorAfter: $cursorAfter) {
    id readableId status takenOverAt sentAt receivedAt processedAt createdAt
    wasteDetails { code name }
    emitter {
      company { siret name address }
      workSite { name address postalCode city }
    }
    recipient {
      processingOperation
      company { siret name address }
    }
    temporaryStorageDetail {
      destination {
        processingOperation
        company { siret name address }
      }
    }
    transporter { company { siret name } }
  }
}
"""


def fetch_all_forms(token, siret, logger=print, page_size=50, max_pages=400):
    headers = {"Authorization": f"Bearer {token.strip()}"}
    all_forms = []
    cursor = None
    page = 0
    while page < max_pages:
        page += 1
        variables = {"siret": siret or None, "first": page_size}
        if cursor:
            variables["cursorAfter"] = cursor
        payload = {"query": FORMS_QUERY, "variables": variables}
        data = http_json(TD_API_URL, method="POST", headers=headers, payload=payload)
        if "errors" in data and data["errors"]:
            msg = "; ".join(e.get("message", "?") for e in data["errors"])
            raise RuntimeError(f"Trackdéchets : {msg}")
        batch = (data.get("data") or {}).get("forms") or []
        if not batch:
            break
        all_forms.extend(batch)
        logger(f"   … {len(all_forms)} BSDD récupérés")
        if len(batch) < page_size:
            break
        cursor = batch[-1]["id"]
    return all_forms


def _addr_of(company):
    if not company:
        return "", ""
    return (company.get("siret") or "").strip(), (company.get("address") or "").strip()


def _emitter_address(form):
    em = form.get("emitter") or {}
    ws = em.get("workSite") or {}
    parts = []
    if ws.get("address"):
        parts.append(ws["address"])
    if ws.get("postalCode"):
        parts.append(ws["postalCode"])
    if ws.get("city"):
        parts.append(ws["city"])
    if parts:
        return " ".join(parts)
    return ((em.get("company") or {}).get("address") or "").strip()


def _year_of(form):
    for k in ("takenOverAt", "sentAt", "receivedAt", "createdAt"):
        v = form.get(k)
        if v:
            try:
                return int(str(v)[:4])
            except Exception:
                continue
    return None


def extract_legs(form, protec_siret):
    protec_siret = (protec_siret or "").replace(" ", "")
    legs = []
    waste = form.get("wasteDetails") or {}
    ced = (waste.get("code") or "").strip()
    year = _year_of(form)

    em = form.get("emitter") or {}
    em_siret, em_addr = _addr_of(em.get("company"))
    em_addr = _emitter_address(form) or em_addr

    rec = form.get("recipient") or {}
    rec_siret, rec_addr = _addr_of(rec.get("company"))
    rec_op = (rec.get("processingOperation") or "").strip()

    tsd = form.get("temporaryStorageDetail") or {}
    dest = (tsd.get("destination") or {}) if tsd else {}
    dest_siret, dest_addr = _addr_of(dest.get("company"))
    dest_op = (dest.get("processingOperation") or "").strip()

    em_norm = em_siret.replace(" ", "")
    rec_norm = rec_siret.replace(" ", "")

    if rec_norm == protec_siret and em_addr:
        legs.append({
            "type": "Collecte", "year": year, "ced": ced,
            "origin_label": (em.get("company") or {}).get("name") or "Producteur",
            "origin_addr": em_addr, "dest_label": "PROTEC", "dest_addr": rec_addr,
            "treatment_code": "",
        })

    if rec_norm == protec_siret and dest_addr:
        legs.append({
            "type": "Reacheminement", "year": year, "ced": ced,
            "origin_label": "PROTEC", "origin_addr": rec_addr,
            "dest_label": (dest.get("company") or {}).get("name") or "Exutoire",
            "dest_addr": dest_addr, "treatment_code": dest_op,
        })

    if em_norm == protec_siret and rec_norm != protec_siret and rec_addr:
        legs.append({
            "type": "Reacheminement", "year": year, "ced": ced,
            "origin_label": "PROTEC", "origin_addr": em_addr,
            "dest_label": (rec.get("company") or {}).get("name") or "Exutoire",
            "dest_addr": rec_addr, "treatment_code": rec_op,
        })

    return legs


def norm_ced(code):
    return "".join(c for c in (code or "") if c.isalnum()).lower()


def categorize(ced, categories):
    n = norm_ced(ced)
    best = None
    best_len = -1
    for cat, prefixes in categories.items():
        for p in prefixes:
            pn = norm_ced(p)
            if pn and n.startswith(pn) and len(pn) > best_len:
                best = cat
                best_len = len(pn)
    return best or "Autres / non classé"


def compute_bilan(forms, cfg, geo, target_years, logger=print, progress=None):
    categories = cfg.get("categories", {})
    protec_siret = cfg.get("protec_siret", "")
    round_trip = bool(cfg.get("round_trip", True))
    factor = 2 if round_trip else 1

    protec_coords = geo.geocode(cfg.get("protec_adresse", "")) if cfg.get("protec_adresse") else None

    rows = []
    summary = {}
    legs_all = []
    for form in forms:
        for leg in extract_legs(form, protec_siret):
            if leg["year"] in target_years:
                legs_all.append((form, leg))

    total = len(legs_all)
    logger(f"   {total} trajet(s) à mesurer pour {sorted(target_years)}")
    done = 0
    for form, leg in legs_all:
        o_addr = leg["origin_addr"]
        d_addr = leg["dest_addr"]
        o_coords = protec_coords if (leg["origin_label"] == "PROTEC" and protec_coords) else geo.geocode(o_addr)
        d_coords = protec_coords if (leg["dest_label"] == "PROTEC" and protec_coords) else geo.geocode(d_addr)

        km1, src = geo.road_km(o_coords, d_coords)
        km = round(km1 * factor, 2) if km1 is not None else None
        cat = categorize(leg["ced"], categories)

        rows.append({
            "BSD": form.get("readableId", ""), "Annee": leg["year"],
            "Type": leg["type"], "Categorie": cat, "CED": leg["ced"],
            "Origine": leg["origin_label"], "Adresse origine": o_addr,
            "Destination": leg["dest_label"], "Adresse destination": d_addr,
            "Code traitement": leg["treatment_code"],
            "Km (1 sens)": km1, "Km comptes": km, "Source": src,
        })

        if km is not None:
            key = (leg["year"], leg["type"], cat, leg["treatment_code"])
            summary[key] = round(summary.get(key, 0.0) + km, 2)

        done += 1
        if progress:
            progress(done, total)
        if done % 25 == 0:
            geo.flush()

    geo.flush()
    return rows, summary


def export_excel(rows, summary, dest, round_trip=True):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("Le module openpyxl est requis : pip install openpyxl")

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F4E5F")
    head_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    def style_header(ws, ncols, row=1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = center
            cell.border = border

    def autosize(ws):
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, 55)

    ws = wb.active
    ws.title = "Synthese"
    note = "ALLER-RETOUR" if round_trip else "aller simple"
    ws.append([f"Bilan kilometrique PROTEC - distances comptees en {note}"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    headers = ["Annee", "Type", "Categorie", "Code traitement", "Km total"]
    ws.append(headers)
    style_header(ws, len(headers), row=3)
    for key in sorted(summary.keys(), key=lambda k: (k[0] or 0, k[1], k[2], k[3])):
        year, typ, cat, code = key
        ws.append([year, typ, cat, code or "-", summary[key]])
    autosize(ws)
    ws.freeze_panes = "A4"

    ws2 = wb.create_sheet("Par_categorie")
    ws2.append(["Annee", "Categorie", "Km Collecte", "Km Reacheminement", "Km Total"])
    style_header(ws2, 5)
    agg = {}
    for (year, typ, cat, code), km in summary.items():
        a = agg.setdefault((year, cat), {"Collecte": 0.0, "Reacheminement": 0.0})
        a[typ] = round(a.get(typ, 0.0) + km, 2)
    for (year, cat) in sorted(agg.keys(), key=lambda k: (k[0] or 0, k[1])):
        c = agg[(year, cat)]["Collecte"]
        r = agg[(year, cat)]["Reacheminement"]
        ws2.append([year, cat, round(c, 2), round(r, 2), round(c + r, 2)])
    autosize(ws2)
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Detail")
    if rows:
        cols = list(rows[0].keys())
        ws3.append(cols)
        style_header(ws3, len(cols))
        for r in rows:
            ws3.append([r.get(c) for c in cols])
        autosize(ws3)
        ws3.freeze_panes = "A2"
    else:
        ws3.append(["Aucun trajet"])

    wb.save(dest)


def parse_years(txt):
    years = set()
    for part in txt.replace(";", ",").split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            years.update(range(int(a.strip()), int(b.strip()) + 1))
        elif part:
            years.add(int(part))
    return years
